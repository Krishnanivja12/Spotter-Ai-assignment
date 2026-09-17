from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_figure(document: Document, path: Path, caption: str) -> None:
    document.add_picture(str(path), width=Inches(6.55))
    paragraph = document.add_paragraph(caption)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.runs[0].italic = True


def main() -> None:
    REPORTS.mkdir(parents=True, exist_ok=True)
    train = pd.read_csv(ROOT / "data" / "train-test.csv")
    validation = pd.read_csv(ROOT / "data" / "validation.csv")
    baseline = json.loads((ROOT / "outputs" / "modeling" / "baseline" / "baseline_metrics.json").read_text(encoding="utf-8"))
    comparison = pd.read_csv(ROOT / "outputs" / "modeling" / "experiments" / "model_comparison.csv")
    manifest = json.loads((ROOT / "outputs" / "modeling" / "final_model_manifest.json").read_text(encoding="utf-8"))

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.color.rgb = RGBColor(6, 74, 86)

    title = document.add_heading("Freight Rate Prediction", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Spotter Machine Learning Engineer Assessment Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Prepared from reproducible project outputs; internal metrics are not official Spotter scores.")

    document.add_heading("1. Executive Summary", 1)
    document.add_paragraph(
        "This project predicts freight rates for unseen loads. The final solution uses a Ridge regression pipeline trained on all "
        "48,000 labeled January–October 2025 loads. It applies deterministic cleaning, compact geographic/calendar features, "
        "training-fitted imputation, and unknown-category-safe encoding. A chronological September–October holdout was used "
        "to simulate future prediction. The selected candidate achieved internal holdout MAE of "
        f"${comparison.iloc[0]['holdout_mae']:.2f}, RMSE of ${comparison.iloc[0]['holdout_rmse']:.2f}, and R² of "
        f"{comparison.iloc[0]['holdout_r2']:.4f}. Final predictions for 12,000 validation loads and 31 December scenario rows "
        "passed the supplied scorer validation."
    )

    document.add_heading("2. Problem Understanding", 1)
    document.add_paragraph(
        "The task is supervised regression: learn posted_rate from labeled historical loads and predict posted_rate for the "
        "unlabeled validation loads. The final validation file contains 12,000 unique load IDs and later dates than development data."
    )

    document.add_heading("3. Dataset Overview", 1)
    add_table(document, ["Dataset", "Rows", "Columns", "Date coverage"], [
        ["Development", "48,000", "14", "2025-01-01 to 2025-10-31"],
        ["Final validation", "12,000", "13", "2025-11-01 to 2025-12-31"],
        ["December scenario", "31", "7", "2025-12-01 to 2025-12-31"],
    ])
    document.add_paragraph("Target: posted_rate. Validation includes the same input schema as development except for the target.")

    document.add_heading("4. Exploratory Data Analysis", 1)
    rate = train["posted_rate"].describe(percentiles=[0.01, 0.05, 0.25, 0.5, 0.75, 0.95, 0.99])
    add_table(document, ["Target statistic", "Value"], [
        ["Minimum", f"${rate['min']:.2f}"], ["Median", f"${rate['50%']:.2f}"],
        ["Mean", f"${rate['mean']:.2f}"], ["Standard deviation", f"${rate['std']:.2f}"],
        ["95th percentile", f"${rate['95%']:.2f}"], ["99th percentile", f"${rate['99%']:.2f}"],
        ["Maximum", f"${rate['max']:.2f}"],
    ])
    document.add_paragraph(
        "Posted rates are right-skewed: the mean exceeds the median and the maximum is substantially above the 99th percentile. "
        "Distance is the strongest numeric relationship with posted_rate (Pearson correlation 0.9085). Reefer loads had the "
        "highest mean rate ($2,553.64), followed by Flatbed ($2,445.09) and Dry Van ($2,271.55)."
    )
    add_figure(document, ROOT / "outputs" / "analysis" / "eda" / "plots" / "03_distance_vs_posted_rate.png", "Figure 1. Distance has a strong positive, but not perfectly linear, relationship with posted rate.")

    document.add_heading("5. Data Quality Issues", 1)
    add_table(document, ["Issue", "Development", "Validation", "Finding"], [
        ["Missing weight", "300", "165", "Requires consistent numeric handling"],
        ["Missing market_index", "374", "249", "Requires consistent numeric handling"],
        ["Negative weight", "292", "145", "Physically invalid"],
        ["Duplicate rows / load IDs", "0 / 0", "0 / 0", "No removal required"],
        ["Validation-only cities", "N/A", "8 pickup and 8 delivery values", "Must support unseen categories"],
    ])
    document.add_paragraph("All dates parsed successfully. No non-positive distances, invalid coordinate ranges, or non-positive training targets were found.")

    document.add_heading("6. Data Cleaning Strategy", 1)
    add_bullets(document, [
        "Negative weights were not dropped or converted to absolute values. They were converted to missing values and retained through weight_was_negative.",
        "Numeric fields are median-imputed inside the training-fitted pipeline, with missing-value indicators.",
        "Categorical fields use most-frequent imputation and one-hot encoding with handle_unknown='ignore'.",
        "No preprocessing statistic is fitted on final validation data.",
    ])

    document.add_heading("7. Feature Engineering", 1)
    document.add_paragraph(
        "The final features retain raw coordinates, distance, weight, market index, quote signal, equipment, and city fields. "
        "Added features are distance_log1p, latitude/longitude differences, haversine distance, distance × market index, "
        "day of week, cyclic day-of-year features, and weight_was_negative. The route string was intentionally rejected: "
        "development had 4,014 routes and 1,461 validation rows used unseen routes."
    )

    document.add_heading("8. Train/Validation Split Strategy", 1)
    add_table(document, ["Partition", "Rows", "Dates", "Purpose"], [
        ["Fit", "38,477", "2025-01-01 to 2025-08-31", "Fit candidate pipelines"],
        ["Internal holdout", "9,523", "2025-09-01 to 2025-10-31", "Future-like model comparison"],
        ["Final validation", "12,000", "2025-11-01 to 2025-12-31", "Unlabeled final prediction"],
    ])
    document.add_paragraph(
        "A chronological cutoff of 2025-09-01 was chosen because the 61-day September–October holdout matches the 61-day "
        "November–December final horizon. A random split would mix earlier and later dates in both partitions and produce a less "
        "realistic estimate of future performance."
    )

    document.add_heading("9. Baseline Model", 1)
    add_table(document, ["Model", "MAE", "RMSE", "R²", "Observation"], [[
        "Ridge baseline", f"${baseline['mae']:.2f}", f"${baseline['rmse']:.2f}", f"{baseline['r2']:.4f}",
        "Underpredicted rare high-rate loads; 11 non-positive raw estimates",
    ]])

    document.add_heading("10. Model Experiments", 1)
    rows = []
    for _, row in comparison.iterrows():
        rows.append([
            str(row["model"]), f"${row['holdout_mae']:.2f}", f"${row['holdout_rmse']:.2f}",
            f"{row['holdout_r2']:.4f}", f"{row['rmse_generalization_gap']:.2f}", f"{row['training_seconds']:.2f}",
        ])
    add_table(document, ["Model", "MAE", "RMSE", "R²", "RMSE gap", "Train sec"], rows)
    document.add_paragraph("All values are internal chronological-holdout metrics, not official Spotter scores.")

    document.add_heading("11. Final Model Selection", 1)
    document.add_paragraph(
        "Ridge was selected because it had the lowest chronological-holdout MAE ($149.63), lowest RMSE ($643.53), highest R² "
        "(0.8222), and smallest RMSE generalization gap (58.44). The tree models fitted historical rows more closely but did not "
        "improve performance on later unseen rows. The final Ridge pipeline includes a 0.01 positive prediction floor required by the scorer."
    )

    document.add_heading("12. Validation Results", 1)
    document.add_paragraph(
        "The selected model's internal evaluation is based only on September–October rows, with all preprocessing fit on January–August. "
        "The project does not claim an official Spotter score: final validation metrics are calculated by Spotter after submission."
    )

    document.add_heading("13. Final Validation Prediction Generation", 1)
    document.add_paragraph(
        "The saved pipeline was retrained on all 48,000 development rows and applied to all 12,000 final validation loads. "
        "validation_predictions.csv contains exactly load_id and predicted_rate, preserves template order, has no duplicate IDs or missing values, "
        "and passed local scorer-compatible checks."
    )

    document.add_heading("14. December Prediction Analysis", 1)
    document.add_paragraph(
        "For the fixed Lexington-to-Fort Wayne scenario (360 miles, Dry Van, 32,000 lb), all 31 December dates were predicted. "
        f"Predictions range from ${manifest['december_inference']['minimum']:.2f} to ${manifest['december_inference']['maximum']:.2f}. "
        "Only date changes across the scenario."
    )

    document.add_heading("15. candidate_december.png", 1)
    add_figure(document, ROOT / "scorer_results" / "candidate_december.png", "Figure 2. Candidate December 2025 predicted load rate, generated by the supplied score.py.")

    document.add_heading("16. Limitations", 1)
    add_bullets(document, [
        "Only one chronological holdout window was used because of the assessment time limit; repeated rolling-origin validation could provide more confidence.",
        "Rare extreme-rate loads remain challenging: Ridge underpredicted the highest actual-rate holdout decile on average.",
        "The final validation period has a lower market_index distribution than development data, creating temporal distribution shift.",
        "December inputs omit market_index and quote_signal; the saved pipeline uses development-fitted imputation for those unavailable fields.",
        "The business provenance of market_index and quote_signal should be confirmed to ensure they are available at quote time.",
    ])

    document.add_heading("17. Reproducibility / How to Run", 1)
    document.add_paragraph("From the repository root:")
    document.add_paragraph(
        "python -m pip install -r requirements.txt\n"
        "python scripts/train_final_model.py\n"
        "python scripts/predict_validation.py\n"
        "python scripts/predict_december.py\n"
        "python score.py --predictions validation_predictions.csv --december-predictions data/december-chart-inputs.csv",
        style="Intense Quote",
    )
    document.add_paragraph("To regenerate this report: python scripts/generate_assessment_report.py")

    output = REPORTS / "report.docx"
    document.save(output)
    print(f"Created report: {output.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
